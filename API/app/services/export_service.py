"""
Servicio para exportar reportes en diferentes formatos (DOCX, XLSX, PDF).
"""

from typing import Optional, Dict, Any, List
from datetime import datetime
from zoneinfo import ZoneInfo
from io import BytesIO
import logging
import re
import textwrap

# Configurar logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Zona horaria de México
MEXICO_TZ = ZoneInfo("America/Mexico_City")


class ExportService:
    """Servicio para exportar reportes en diferentes formatos."""

    def __init__(self):
        """Inicializar el servicio de exportación."""
        pass

    def export_to_word(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> BytesIO:
        """
        Exportar reporte a formato Word (DOCX).

        Args:
            contenido: Contenido del reporte en Markdown
            metadata: Metadatos del reporte (versión, fecha, etc.)

        Returns:
            BytesIO: Archivo Word en memoria
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            # Crear documento
            doc = Document()

            # Agregar encabezado con metadatos
            if metadata:
                header = doc.sections[0].header
                header_para = header.paragraphs[0]
                header_para.text = f"Reporte Psicopedagógico - Versión {metadata.get('version', 'N/A')}"
                header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                
                # Agregar información de fecha
                if metadata.get('fecha_generacion'):
                    fecha_para = doc.add_paragraph()
                    fecha_para.add_run(f"Fecha de generación: {metadata['fecha_generacion']}").italic = True
                    fecha_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            # Agregar línea separadora
            doc.add_paragraph('_' * 80)

            # Procesar contenido Markdown y convertir a Word
            self._markdown_to_word(doc, contenido)

            # Agregar pie de página
            footer = doc.sections[0].footer
            footer_para = footer.paragraphs[0]
            footer_para.text = f"Generado el {datetime.now(MEXICO_TZ).strftime('%d/%m/%Y %H:%M')}"
            footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Guardar en BytesIO
            file_stream = BytesIO()
            doc.save(file_stream)
            file_stream.seek(0)

            return file_stream

        except Exception as e:
            logger.error(f"Error exportando a Word: {e}")
            raise

    def export_to_excel(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> BytesIO:
        """
        Exportar reporte a formato Excel (XLSX).

        Args:
            contenido: Contenido del reporte en Markdown
            metadata: Metadatos del reporte (versión, fecha, etc.)

        Returns:
            BytesIO: Archivo Excel en memoria
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter

            # Crear workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte Psicopedagógico"

            # Configurar anchos de columna
            ws.column_dimensions['A'].width = 100

            row = 1

            # Agregar encabezado con metadatos
            if metadata:
                ws.cell(row=row, column=1, value="REPORTE PSICOPEDAGÓGICO")
                ws.cell(row=row, column=1).font = Font(size=16, bold=True)
                ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
                row += 1

                ws.cell(row=row, column=1, value=f"Versión: {metadata.get('version', 'N/A')}")
                ws.cell(row=row, column=1).font = Font(size=12, italic=True)
                row += 1

                if metadata.get('fecha_generacion'):
                    ws.cell(row=row, column=1, value=f"Fecha de generación: {metadata['fecha_generacion']}")
                    ws.cell(row=row, column=1).font = Font(size=10, italic=True)
                    row += 1

                row += 1  # Línea en blanco

            # Procesar contenido Markdown
            lines = contenido.split('\n')
            for line in lines:
                if not line.strip():
                    row += 1
                    continue

                cell = ws.cell(row=row, column=1, value=line)
                
                # Aplicar estilos según el tipo de línea
                if line.startswith('# '):
                    cell.font = Font(size=14, bold=True, color="1F4E78")
                    cell.value = line.replace('# ', '')
                elif line.startswith('## '):
                    cell.font = Font(size=12, bold=True, color="2E5C8A")
                    cell.value = line.replace('## ', '')
                elif line.startswith('### '):
                    cell.font = Font(size=11, bold=True, color="4472C4")
                    cell.value = line.replace('### ', '')
                elif line.startswith('**') and line.endswith('**'):
                    cell.font = Font(bold=True)
                    cell.value = line.replace('**', '')
                elif 'URGENTE' in line:
                    cell.fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
                    cell.font = Font(bold=True, color="856404")

                # Ajustar altura de fila
                ws.row_dimensions[row].height = 15 if len(line) < 100 else 30
                
                row += 1

            # Guardar en BytesIO
            file_stream = BytesIO()
            wb.save(file_stream)
            file_stream.seek(0)

            return file_stream

        except Exception as e:
            logger.error(f"Error exportando a Excel: {e}")
            raise

    def export_to_pdf(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> BytesIO:
        """
        Exportar reporte a formato PDF.

        Args:
            contenido: Contenido del reporte en Markdown
            metadata: Metadatos del reporte (versión, fecha, etc.)

        Returns:
            BytesIO: Archivo PDF en memoria
        """
        try:
            from weasyprint import HTML, CSS
            from app.services.markdown_service import markdown_service

            # Convertir Markdown a HTML
            html_content = markdown_service.render_to_html(contenido, add_css=True)

            # Agregar metadatos al HTML
            if metadata:
                metadata_html = f"""
                <div style="text-align: right; margin-bottom: 20px; padding: 10px; background-color: #f0f0f0; border-radius: 5px;">
                    <p style="margin: 5px 0;"><strong>Versión:</strong> {metadata.get('version', 'N/A')}</p>
                    <p style="margin: 5px 0;"><strong>Fecha de generación:</strong> {metadata.get('fecha_generacion', 'N/A')}</p>
                </div>
                """
                # Insertar metadatos después del <body>
                html_content = html_content.replace('<body>', f'<body>{metadata_html}')

            # CSS adicional para PDF
            pdf_css = CSS(string="""
                @page {
                    size: A4;
                    margin: 2cm;
                }
                body {
                    font-size: 11pt;
                }
                h1 {
                    page-break-before: avoid;
                }
            """)

            # Generar PDF
            pdf_file = HTML(string=html_content).write_pdf(stylesheets=[pdf_css])

            # Crear BytesIO
            file_stream = BytesIO(pdf_file)
            file_stream.seek(0)

            return file_stream

        except (ImportError, OSError) as e:
            logger.warning(
                "Dependencias de WeasyPrint no disponibles, usando exportación PDF simplificada: %s",
                e
            )
            return self._export_to_pdf_simple(contenido, metadata)
        except Exception as e:
            logger.error(f"Error exportando a PDF con WeasyPrint: {e}")
            # Intentar fallback simplificado
            try:
                return self._export_to_pdf_simple(contenido, metadata)
            except Exception as fallback_error:
                logger.error(f"Error en exportación PDF simplificada: {fallback_error}")
                raise

    def _export_to_pdf_simple(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> BytesIO:
        """Generar un PDF usando ReportLab si está disponible; de lo contrario, modo básico."""
        try:
            return self._generate_reportlab_pdf(contenido, metadata)
        except ImportError:
            logger.warning("ReportLab no está instalado. Generando PDF básico.")
        except Exception as e:
            logger.warning("Error exportando PDF con ReportLab: %s. Usando modo básico.", e)

        plain_text = self._build_plain_text(contenido, metadata)
        return self._generate_simple_pdf(plain_text)

    def _generate_reportlab_pdf(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> BytesIO:
        from xml.sax.saxutils import escape

        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            ListFlowable,
            ListItem
        )

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=0.75 * inch,
            rightMargin=0.75 * inch,
            topMargin=1 * inch,
            bottomMargin=1 * inch
        )

        styles = getSampleStyleSheet()
        body_style = ParagraphStyle(
            'BodyTextCustom',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=11,
            leading=14
        )
        heading1 = ParagraphStyle('Heading1Custom', parent=styles['Heading1'], fontName='Helvetica-Bold')
        heading2 = ParagraphStyle('Heading2Custom', parent=styles['Heading2'], fontName='Helvetica-Bold')
        heading3 = ParagraphStyle('Heading3Custom', parent=styles['Heading3'], fontName='Helvetica-Bold')
        bullet_style = ParagraphStyle(
            'BulletCustom',
            parent=body_style,
            leftIndent=18,
            bulletIndent=9,
            leading=13
        )
        numbered_style = ParagraphStyle('NumberedCustom', parent=body_style, leftIndent=18)
        metadata_style = ParagraphStyle(
            'Metadata',
            parent=body_style,
            fontSize=9,
            leading=11,
            alignment=TA_RIGHT
        )
        code_style = ParagraphStyle(
            'CodeBlock',
            parent=body_style,
            fontName='Courier',
            fontSize=9,
            leading=11,
            backColor='#f5f5f5',
            leftIndent=12,
            rightIndent=12,
            spaceBefore=6,
            spaceAfter=6
        )

        story = []

        if metadata:
            meta_lines = []
            if metadata.get('version') is not None:
                meta_lines.append(f"<b>Versión:</b> {escape(str(metadata.get('version', 'N/A')))}")
            if metadata.get('fecha_generacion') is not None:
                meta_lines.append(f"<b>Fecha de generación:</b> {escape(str(metadata.get('fecha_generacion', 'N/A')))}")
            if meta_lines:
                story.append(Paragraph('<br/>'.join(meta_lines), metadata_style))
                story.append(Spacer(1, 12))

        def preprocess_links(text: str) -> str:
            def replace_link(match: re.Match) -> str:
                label, url = match.group(1), match.group(2)
                return f"{label} ({url})"
            return re.sub(r'\[([^\]]+)\]\(([^\)]+)\)', replace_link, text)

        def format_inline(text: str) -> str:
            text = preprocess_links(text)
            text = escape(text)
            text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
            text = re.sub(r'__(.+?)__', r'<b>\1</b>', text)
            text = re.sub(r'\*(.+?)\*', r'<i>\1</i>', text)
            text = re.sub(r'_(.+?)_', r'<i>\1</i>', text)
            text = re.sub(r'`(.+?)`', r'<font face="Courier">\1</font>', text)
            return text

        def flush_code_block(buffered_lines: List[str]) -> None:
            if not buffered_lines:
                return
            code_html = '<br/>'.join(escape(line) for line in buffered_lines)
            story.append(Paragraph(f'<font face="Courier">{code_html}</font>', code_style))
            story.append(Spacer(1, 6))
            buffered_lines.clear()

        current_bullet_items: List[ListItem] = []

        def flush_bullet_list() -> None:
            nonlocal current_bullet_items
            if not current_bullet_items:
                return
            story.append(ListFlowable(current_bullet_items, bulletType='bullet'))
            story.append(Spacer(1, 6))
            current_bullet_items = []

        def add_bullet_item(text: str) -> None:
            para = Paragraph(format_inline(text), bullet_style)
            current_bullet_items.append(ListItem(para, bulletText='•'))

        in_code_block = False
        buffered_code: List[str] = []

        lines = contenido.splitlines()
        for line in lines:
            stripped = line.strip()

            if stripped.startswith('```'):
                if in_code_block:
                    flush_code_block(buffered_code)
                    in_code_block = False
                else:
                    flush_bullet_list()
                    in_code_block = True
                    buffered_code = []
                continue

            if in_code_block:
                buffered_code.append(line.rstrip('\n'))
                continue

            if not stripped:
                flush_bullet_list()
                story.append(Spacer(1, 6))
                continue

            if stripped.startswith('# '):
                flush_bullet_list()
                story.append(Paragraph(format_inline(stripped[2:]), heading1))
                story.append(Spacer(1, 6))
                continue

            if stripped.startswith('## '):
                flush_bullet_list()
                story.append(Paragraph(format_inline(stripped[3:]), heading2))
                story.append(Spacer(1, 4))
                continue

            if stripped.startswith('### '):
                flush_bullet_list()
                story.append(Paragraph(format_inline(stripped[4:]), heading3))
                story.append(Spacer(1, 4))
                continue

            if stripped.startswith(('- ', '* ')):
                add_bullet_item(stripped[2:].strip())
                continue

            ordered_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
            if ordered_match:
                flush_bullet_list()
                number = escape(ordered_match.group(1))
                text = format_inline(ordered_match.group(2))
                story.append(Paragraph(f"{number}. {text}", numbered_style))
                continue

            flush_bullet_list()
            story.append(Paragraph(format_inline(stripped), body_style))

        flush_code_block(buffered_code)
        flush_bullet_list()

        doc.build(story)
        buffer.seek(0)
        return buffer

    def _build_plain_text(
        self,
        contenido: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[str]:
        """Construir líneas de texto plano a partir del contenido Markdown y metadatos."""
        lines: List[str] = []

        if metadata:
            version = metadata.get('version', 'N/A')
            fecha = metadata.get('fecha_generacion', 'N/A')
            lines.append("Reporte Psicopedagógico")
            lines.append(f"Versión: {version}")
            lines.append(f"Fecha de generación: {fecha}")
            lines.append('')

        stripped = self._strip_markdown(contenido)
        for raw_line in stripped.splitlines():
            wrapped = textwrap.wrap(raw_line, width=90) or ['']
            lines.extend(wrapped)

        return lines

    def _strip_markdown(self, text: str) -> str:
        """Eliminar formateo Markdown para obtener texto plano."""
        # Eliminar bloques de código
        text = re.sub(r'```[\s\S]*?```', '', text)
        # Eliminar imágenes
        text = re.sub(r'!\[[^\]]*\]\([^\)]*\)', '', text)
        # Convertir vínculos [texto](url) -> texto
        text = re.sub(r'\[([^\]]+)\]\(([^\)]+)\)', r'\1', text)
        # Encabezados
        text = re.sub(r'^#{1,6}\s*', '', text, flags=re.MULTILINE)
        # Negritas y cursivas
        text = text.replace('**', '').replace('__', '')
        text = text.replace('*', '').replace('_', '')
        # Código inline
        text = text.replace('`', '')
        # Ajustar espacios múltiples
        text = re.sub(r'\s+\n', '\n', text)
        return text

    def _generate_simple_pdf(self, lines: List[str]) -> BytesIO:
        """Crear un PDF muy básico con texto plano a partir de una lista de líneas."""
        if not lines:
            lines = ['(Sin contenido)']

        page_width = 612  # 8.5 pulgadas
        page_height = 792  # 11 pulgadas
        margin_left = 50
        margin_top = 50
        margin_bottom = 50
        line_height = 14

        available_height = page_height - margin_top - margin_bottom
        max_lines_per_page = max(1, available_height // line_height)

        pages: List[List[str]] = []
        current_page: List[str] = []

        for line in lines:
            if len(current_page) >= max_lines_per_page:
                pages.append(current_page)
                current_page = []
            current_page.append(line)

        if current_page:
            pages.append(current_page)

        if not pages:
            pages = [['(Sin contenido)']]

        # Construir contenidos de página
        page_contents: List[str] = []
        def to_pdf_hex(text: str) -> str:
            utf16 = text.encode('utf-16-be')
            return 'FEFF' + utf16.hex()

        for page_lines in pages:
            y_position = page_height - margin_top
            content_lines = ['BT', '/F1 12 Tf']
            for line in page_lines:
                hex_text = to_pdf_hex(line)
                content_lines.append(f"1 0 0 1 {margin_left} {y_position} Tm <{hex_text}> Tj")
                y_position -= line_height
            content_lines.append('ET')
            page_contents.append('\n'.join(content_lines))

        num_pages = len(page_contents)
        font_id = 1
        content_ids = [i + 2 for i in range(num_pages)]
        first_page_id = 2 + num_pages
        page_ids = [first_page_id + i for i in range(num_pages)]
        pages_id = first_page_id + num_pages
        catalog_id = pages_id + 1

        objects: List[str] = []

        # 1. Fuente
        objects.append('<< /Type /Font /Subtype /Type1 /Name /F1 /BaseFont /Helvetica >>')

        # 2. Contenidos de página
        for content in page_contents:
            content_bytes = content.encode('latin-1', 'replace')
            stream = f"<< /Length {len(content_bytes)} >>\nstream\n" + content_bytes.decode('latin-1', 'replace') + "\nendstream"
            objects.append(stream)

        # 3. Objetos de página (se completan después de conocer pages_id)
        for idx, content_id in enumerate(content_ids):
            page_obj = (
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {page_width} {page_height}] "
                f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
            )
            objects.append(page_obj)

        # 4. Objeto /Pages
        kids = ' '.join(f"{pid} 0 R" for pid in page_ids)
        pages_obj = f"<< /Type /Pages /Kids [{kids}] /Count {num_pages} >>"
        objects.append(pages_obj)

        # 5. Catálogo
        catalog_obj = f"<< /Type /Catalog /Pages {pages_id} 0 R >>"
        objects.append(catalog_obj)

        # Generar PDF binario
        buffer = BytesIO()
        buffer.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")

        offsets = []
        for index, obj in enumerate(objects, start=1):
            offsets.append(buffer.tell())
            buffer.write(f"{index} 0 obj\n".encode('latin-1'))
            buffer.write(obj.encode('latin-1', 'replace'))
            buffer.write(b"\nendobj\n")

        xref_position = buffer.tell()
        buffer.write(f"xref\n0 {len(objects)+1}\n".encode('latin-1'))
        buffer.write(b"0000000000 65535 f \n")
        for offset in offsets:
            buffer.write(f"{offset:010d} 00000 n \n".encode('latin-1'))

        buffer.write(b"trailer\n")
        buffer.write(f"<< /Size {len(objects)+1} /Root {catalog_id} 0 R >>\n".encode('latin-1'))
        buffer.write(b"startxref\n")
        buffer.write(f"{xref_position}\n".encode('latin-1'))
        buffer.write(b"%%EOF")
        buffer.seek(0)

        return buffer

    def _markdown_to_word(self, doc, markdown_text: str):
        """
        Convertir texto Markdown a formato Word.

        Args:
            doc: Documento de Word
            markdown_text: Texto en formato Markdown
        """
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        lines = markdown_text.split('\n')
        
        for line in lines:
            if not line.strip():
                doc.add_paragraph()
                continue

            # Títulos nivel 1
            if line.startswith('# '):
                para = doc.add_heading(line.replace('# ', ''), level=1)
            # Títulos nivel 2
            elif line.startswith('## '):
                para = doc.add_heading(line.replace('## ', ''), level=2)
            # Títulos nivel 3
            elif line.startswith('### '):
                para = doc.add_heading(line.replace('### ', ''), level=3)
            # Listas con viñetas
            elif line.strip().startswith('- ') or line.strip().startswith('* '):
                para = doc.add_paragraph(line.strip()[2:], style='List Bullet')
            # Listas numeradas
            elif re.match(r'^\d+\.\s', line.strip()):
                para = doc.add_paragraph(re.sub(r'^\d+\.\s', '', line.strip()), style='List Number')
            # Texto normal
            else:
                para = doc.add_paragraph(line)
                
                # Resaltar palabras clave
                if 'URGENTE' in line:
                    for run in para.runs:
                        if 'URGENTE' in run.text:
                            run.font.color.rgb = RGBColor(220, 53, 69)
                            run.font.bold = True


# Instancia global del servicio
export_service = ExportService()
